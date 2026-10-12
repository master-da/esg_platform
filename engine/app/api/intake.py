"""Category pages: templates, uploads, file status & actions, Data Overview (PRD F4, F5, F5.1, F5.2)."""

from __future__ import annotations

import io
from typing import Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app.api.projects import get_project
from app.auth import require_internal, require_read, require_write, user_tx
from app.db import Tx
from app.ingest.service import handle_upload
from app.ingest.templates import build_template
from app.ingest.terminal import Terminal, load_terminal
from app.supabase_admin import storage_signed_url

router = APIRouter(tags=["intake"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(data: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(data),
        media_type=XLSX,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def get_terminal(tx: Tx, key: str) -> Terminal:
    t = load_terminal(tx, key)
    if not t:
        raise HTTPException(404, "Unknown category")
    return t


def get_file(tx: Tx, file_id: str) -> dict:
    f = tx.one("select * from public.uploaded_files where id = %s and deleted_at is null", (file_id,))
    if not f:
        raise HTTPException(404, "File not found")
    return f


@router.get("/projects/{project_id}/terminals/{key}")
def terminal_page(project_id: str, key: str, tx: Tx = Depends(user_tx)):
    p = get_project(tx, project_id)
    t = get_terminal(tx, key)
    pc = tx.one(
        """select pc.*, pr.full_name as assignee_name from public.project_categories pc
           left join public.profiles pr on pr.user_id = pc.assignee_id
           where pc.project_id = %s and pc.terminal_key = %s""",
        (project_id, key),
    )
    files = tx.all(
        """select f.*, pr.full_name as uploaded_by_name from public.v_file_status f
           left join public.profiles pr on pr.user_id = f.uploaded_by
           where f.project_id = %s and f.terminal_key = %s order by f.uploaded_at desc""",
        (project_id, key),
    )
    needs = tx.scalar(
        """select count(*) from public.categorizations c where c.project_id = %s and c.terminal_key = %s
             and c.status <> 'ACCEPTED' and exists (select 1 from public.activity_records r where r.project_id = c.project_id
             and r.terminal_key = c.terminal_key and r.category_input_key = c.input_key and r.deleted_at is null)""",
        (project_id, key),
    )
    return {
        "project": p,
        "terminal": {
            "key": t.key,
            "name": t.name,
            "label": t.label,
            "scope": t.scope,
            "category_no": t.category_no,
            "module": t.module,
            "description": t.description,
            "excludes": t.excludes,
            "instructions": t.instructions,
            "accepts_bills": t.accepts_bills,
            "enabled": t.enabled,
            "disabled_reason": t.disabled_reason,
            "template_version": t.template_version,
            "columns": [
                {"key": c.key, "header": c.header, "type": c.type, "required": c.required, "options": list(c.options)}
                for c in t.columns
            ],
            "categorization_keys": list(t.categorization_keys),
            "bill_fields": t.bill_fields,
        },
        "category": pc,
        "files": files,
        "needs_categorization": needs,
        "need_attention_files": sum(1 for f in files if f["status"] == "NEED_ATTENTION"),
    }


@router.get("/projects/{project_id}/terminals/{key}/template")
def download_template(project_id: str, key: str, tx: Tx = Depends(user_tx)):
    p = get_project(tx, project_id)
    t = get_terminal(tx, key)
    facilities = [
        r["name"]
        for r in tx.all(
            "select name from public.facilities where company_id = %s order by is_company_level, name",
            (p["company_id"],),
        )
    ]
    units = {r["code"]: r["name"] for r in tx.all("select code, name from ref.units")}
    return xlsx_response(build_template(t, facilities, units), f"{t.key}_template_v{t.template_version}.xlsx")


@router.post("/projects/{project_id}/terminals/{key}/uploads", status_code=201)
def upload(project_id: str, key: str, file: UploadFile = File(...), tx: Tx = Depends(user_tx)):
    p = get_project(tx, project_id)
    require_write(tx, str(p["company_id"]))
    t = get_terminal(tx, key)
    return handle_upload(tx, p, t, file.filename or "upload", file.file.read(), file.content_type)


@router.get("/files/{file_id}/issues")
def file_issues(file_id: str, tx: Tx = Depends(user_tx)):
    get_file(tx, file_id)
    return tx.all(
        """select row_number, field, problem, fix, severity from public.validation_issues where file_id = %s
                     order by row_number nulls first, field""",
        (file_id,),
    )


@router.get("/files/{file_id}/issues.xlsx")
def file_issue_report(file_id: str, tx: Tx = Depends(user_tx)):
    f = get_file(tx, file_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append(["Row", "Field", "Problem", "Fix"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i in file_issues(file_id, tx):
        ws.append([i["row_number"], i["field"], i["problem"], i["fix"]])
    buf = io.BytesIO()
    wb.save(buf)
    return xlsx_response(buf.getvalue(), f"errors_{f['original_name'].rsplit('.', 1)[0]}.xlsx")


@router.get("/files/{file_id}/original")
def download_original(file_id: str, tx: Tx = Depends(user_tx)):
    """All users who can see the file may download the original (signed URL, short-lived)."""
    f = get_file(tx, file_id)
    tx.log_event(str(f["company_id"]), "public.uploaded_files", file_id, "DOWNLOAD_ORIGINAL", None)
    return {"url": storage_signed_url(f["storage_path"], 120, f["original_name"])}


@router.get("/files/{file_id}/calculated")
def download_calculated(file_id: str, tx: Tx = Depends(user_tx)):
    """Internal users only (PRD 5.1): the uploaded rows with their calculated results."""
    require_internal(tx)
    f = get_file(tx, file_id)
    rows = tx.all(
        """select r.source_row, fa.name as facility, r.supplier_name, r.period_start, r.period_end,
                  r.original_quantity, r.original_unit, r.quality_grade, cc.status, cc.taxonomy_display,
                  cc.conversion_from, cc.conversion_to, cc.conversion_multiplier, cc.factor_value, cc.factor_unit,
                  cc.factor_source_name, cc.factor_source_year, cc.factor_version, cc.co2e_kg, cc.co2e_t, cc.formula
           from public.activity_records r
           join public.facilities fa on fa.id = r.facility_id
           left join public.v_current_calcs cc on cc.activity_record_id = r.id
           where r.file_id = %s and r.deleted_at is null order by r.source_row""",
        (file_id,),
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Calculated"
    headers = [
        "Source row",
        "Facility",
        "Supplier",
        "Period start",
        "Period end",
        "Input",
        "Input unit",
        "Grade",
        "Status",
        "Category",
        "Conversion from",
        "Conversion to",
        "Conversion factor",
        "Emission factor",
        "Factor unit (per)",
        "Factor database",
        "Factor year",
        "Factor version",
        "kg CO2e",
        "t CO2e",
        "Formula",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                r["source_row"],
                r["facility"],
                r["supplier_name"],
                r["period_start"],
                r["period_end"],
                r["original_quantity"],
                r["original_unit"],
                r["quality_grade"],
                r["status"] or "NOT CATEGORIZED",
                r["taxonomy_display"],
                r["conversion_from"],
                r["conversion_to"],
                r["conversion_multiplier"],
                r["factor_value"],
                r["factor_unit"],
                r["factor_source_name"],
                r["factor_source_year"],
                r["factor_version"],
                r["co2e_kg"],
                r["co2e_t"],
                r["formula"],
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    tx.log_event(str(f["company_id"]), "public.uploaded_files", file_id, "DOWNLOAD_CALCULATED", None)
    return xlsx_response(buf.getvalue(), f"calculated_{f['original_name'].rsplit('.', 1)[0]}.xlsx")


@router.delete("/files/{file_id}")
def delete_file(file_id: str, tx: Tx = Depends(user_tx)):
    """Removes the file and its records from the active dataset; the audit log keeps the history."""
    f = get_file(tx, file_id)
    require_write(tx, str(f["company_id"]))
    with tx.privileged():
        tx.execute(
            "update public.activity_records set deleted_at = now(), deleted_by = app.uid() "
            "where file_id = %s and deleted_at is null",
            (file_id,),
        )
        tx.execute(
            "update public.uploaded_files set deleted_at = now(), deleted_by = app.uid() where id = %s", (file_id,)
        )
    tx.log_event(
        str(f["company_id"]),
        "public.uploaded_files",
        file_id,
        "DELETE_FILE",
        {"file": f["original_name"], "terminal": f["terminal_key"]},
    )
    return {"deleted": True}


class ReviewIn(BaseModel):
    action: Literal["confirm", "reject"]
    comment: str | None = None


@router.post("/files/{file_id}/review")
def review_file(file_id: str, body: ReviewIn, tx: Tx = Depends(user_tx)):
    """Internal reviewers confirm a file as Calculated, or return it to the client (PRD 5.1)."""
    require_internal(tx)
    f = get_file(tx, file_id)
    status = tx.one("select status from public.v_file_status where file_id = %s", (file_id,))["status"]
    if body.action == "confirm" and status != "PRELIMINARY":
        raise HTTPException(409, f"Only files with Preliminary calculations can be confirmed (this one is {status}).")
    if body.action == "reject" and not (body.comment or "").strip():
        raise HTTPException(422, "Say what the client needs to fix when returning a file")
    p = get_project(tx, str(f["project_id"]))
    if p["stage"] in ("APPROVED", "LOCKED"):
        raise HTTPException(409, f"The project is {p['stage']}; files can no longer be reviewed.")
    with tx.privileged():
        tx.execute(
            """insert into public.file_reviews (company_id, file_id, action, comment, calc_version_id)
                      values (%s, %s, %s, %s, %s)""",
            (f["company_id"], file_id, body.action, body.comment, p["current_version_id"]),
        )
    tx.log_event(
        str(f["company_id"]),
        "public.uploaded_files",
        file_id,
        "CONFIRM_CALCULATED" if body.action == "confirm" else "REJECT",
        {"comment": body.comment},
    )
    return {"reviewed": body.action}


@router.get("/projects/{project_id}/terminals/{key}/records")
def records(project_id: str, key: str, tx: Tx = Depends(user_tx)):
    """Data Overview: every record with input and calculated values (PRD F5)."""
    p = get_project(tx, project_id)
    require_read(tx, str(p["company_id"]))
    return tx.all(
        """select r.id, r.file_id, uf.original_name as file_name, r.source_row, r.source_type, fa.name as facility,
                  r.facility_id, r.supplier_name, r.period_start, r.period_end, r.attributes, r.category_input,
                  r.original_quantity, r.original_unit, r.normalized_quantity, r.normalized_unit, r.quality_grade,
                  r.review_status, c.status as categorization_status, l.display_path as categorized_as,
                  cc.id as calculation_id, cc.status as calc_status, cc.message as calc_message, cc.co2e_t,
                  fs.status as file_status
           from public.activity_records r
           join public.facilities fa on fa.id = r.facility_id
           left join public.uploaded_files uf on uf.id = r.file_id
           left join public.v_file_status fs on fs.file_id = r.file_id
           left join public.categorizations c on c.project_id = r.project_id and c.terminal_key = r.terminal_key
                and c.input_key = r.category_input_key
           left join ref.taxonomy_labels l on l.id = c.accepted_node_id
           left join public.v_current_calcs cc on cc.activity_record_id = r.id
           where r.project_id = %s and r.terminal_key = %s and r.deleted_at is null
           order by uf.uploaded_at desc nulls last, r.source_row""",
        (project_id, key),
    )
