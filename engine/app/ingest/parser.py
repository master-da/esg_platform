"""Deterministic parsing of uploaded template files (.xlsx or .csv). No AI is involved."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from app.ingest.templates import DATA_SHEET, META_SHEET
from app.ingest.terminal import Terminal, norm


class ParseError(Exception):
    """The file cannot be read at all (wrong format, empty, corrupt)."""


@dataclass
class ParsedFile:
    meta: dict[str, str]
    headers: list[str]
    rows: list[tuple[int, dict[str, Any]]] = field(default_factory=list)  # (spreadsheet row number, header -> value)


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def parse_file(filename: str, data: bytes, max_rows: int) -> ParsedFile:
    name = filename.lower()
    if name.endswith(".xlsx"):
        return _parse_xlsx(data, max_rows)
    if name.endswith(".csv"):
        return _parse_csv(data, max_rows)
    if name.endswith(".xls"):
        raise ParseError("Legacy .xls files are not supported. Save the file as .xlsx and upload it again.")
    raise ParseError("Unsupported file type. Upload this category's .xlsx template (or a .csv with the same columns).")


def _parse_xlsx(data: bytes, max_rows: int) -> ParsedFile:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for corrupt files
        raise ParseError("The file could not be opened as an Excel workbook.") from exc
    meta: dict[str, str] = {}
    if META_SHEET in wb.sheetnames:
        for row in wb[META_SHEET].iter_rows(min_row=1, max_col=2, values_only=True):
            if row and row[0]:
                meta[str(row[0])] = "" if row[1] is None else str(row[1])
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header_row = next(it)
    except StopIteration as exc:
        raise ParseError("The file is empty.") from exc
    headers = ["" if h is None else str(h) for h in header_row]
    parsed = ParsedFile(meta=meta, headers=headers)
    for n, values in enumerate(it, start=2):
        if all(_blank(v) for v in values):
            continue
        parsed.rows.append((n, {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}))
        if len(parsed.rows) > max_rows:
            raise ParseError(f"The file has more than {max_rows} data rows. Split it into smaller files.")
    return parsed


def _parse_csv(data: bytes, max_rows: int) -> ParsedFile:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseError("CSV files must be UTF-8 encoded.") from exc
    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise ParseError("The file is empty.") from exc
    parsed = ParsedFile(meta={}, headers=headers)
    for n, values in enumerate(reader, start=2):
        if all(_blank(v) for v in values):
            continue
        parsed.rows.append((n, {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}))
        if len(parsed.rows) > max_rows:
            raise ParseError(f"The file has more than {max_rows} data rows. Split it into smaller files.")
    return parsed


def detect_terminal(parsed: ParsedFile, terminals: dict[str, Terminal]) -> str | None:
    """Which terminal's template is this? Uses the hidden _meta sheet, else the header signature."""
    key = parsed.meta.get("terminal_key")
    if key in terminals:
        return key
    got = {norm(h) for h in parsed.headers if h}
    best, best_score = None, 0.0
    for t in terminals.values():
        want = {norm(c.header) for c in t.columns}
        if not want:
            continue
        score = len(got & want) / len(got | want)
        if score > best_score:
            best, best_score = t.key, score
    return best if best_score >= 0.8 else None
