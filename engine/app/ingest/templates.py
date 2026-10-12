"""Excel template generation per upload terminal (PRD F5, Method 1).

Each template has a "Data" sheet with dropdown-validated columns, an "Instructions" sheet, a hidden
"_lists" sheet holding the dropdown values, and a hidden "_meta" sheet that identifies the
terminal, so a file uploaded to the wrong terminal is recognised and rejected.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.ingest.terminal import Terminal

DATA_ROWS = 1000
HEADER_FILL = PatternFill("solid", fgColor="1F5C45")
REQUIRED_FILL = PatternFill("solid", fgColor="2F7A5D")
META_SHEET = "_meta"
LISTS_SHEET = "_lists"
DATA_SHEET = "Data"


def build_template(terminal: Terminal, facilities: list[str], unit_names: dict[str, str]) -> bytes:
    wb = Workbook()
    data = wb.active
    data.title = DATA_SHEET
    lists = wb.create_sheet(LISTS_SHEET)
    meta = wb.create_sheet(META_SHEET)

    list_col = 0
    for idx, col in enumerate(terminal.columns, start=1):
        letter = get_column_letter(idx)
        cell = data.cell(row=1, column=idx, value=col.header + (" *" if col.required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if col.required else HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        data.column_dimensions[letter].width = max(14, min(36, len(col.header) + 6))
        rng = f"{letter}2:{letter}{DATA_ROWS + 1}"

        options: list[str] | None = None
        if col.type == "facility":
            options = facilities
        elif col.type in ("enum", "unit"):
            options = list(col.options)
        if options:
            list_col += 1
            lc = get_column_letter(list_col)
            lists.cell(row=1, column=list_col, value=col.header)
            for i, opt in enumerate(options, start=2):
                lists.cell(row=i, column=list_col, value=opt)
            dv = DataValidation(
                type="list",
                formula1=f"={LISTS_SHEET}!${lc}$2:${lc}${len(options) + 1}",
                allow_blank=not col.required,
                showErrorMessage=True,
                errorTitle="Not in list",
                error=f"Choose a value from the {col.header} list",
            )
            data.add_data_validation(dv)
            dv.add(rng)
        elif col.type == "date":
            dv = DataValidation(
                type="date",
                operator="greaterThan",
                formula1="DATE(2000,1,1)",
                allow_blank=not col.required,
                showErrorMessage=True,
                error="Enter a date (YYYY-MM-DD)",
            )
            data.add_data_validation(dv)
            dv.add(rng)
            for r in range(2, DATA_ROWS + 2):
                data.cell(row=r, column=idx).number_format = "yyyy-mm-dd"
        elif col.type in ("decimal", "integer"):
            dv = DataValidation(
                type="decimal" if col.type == "decimal" else "whole",
                operator="greaterThanOrEqual",
                formula1=str(col.min if col.min is not None else -1e15),
                allow_blank=not col.required,
                showErrorMessage=True,
                error="Enter a number" + (f" ≥ {col.min:g}" if col.min is not None else ""),
            )
            data.add_data_validation(dv)
            dv.add(rng)
    data.freeze_panes = "A2"
    data.row_dimensions[1].height = 30

    ins = wb.create_sheet("Instructions", 1)
    ins.column_dimensions["A"].width = 28
    ins.column_dimensions["B"].width = 90
    rows = [
        (terminal.label, None),
        ("What belongs here", terminal.description),
        ("What does not", terminal.excludes or "—"),
        ("How to fill it in", terminal.instructions or "—"),
        ("Template version", terminal.template_version),
        (None, None),
        ("Column", "Guidance"),
    ]
    for col in terminal.columns:
        guide = []
        guide.append("Required." if col.required else "Optional.")
        if col.type == "date":
            guide.append("Date as YYYY-MM-DD.")
        if col.type in ("decimal", "integer"):
            guide.append("Number" + (f", at least {col.min:g}." if col.min is not None else "."))
        if col.type == "facility":
            guide.append("Choose one of the company's facilities.")
        if col.type == "unit":
            guide.append("Units: " + ", ".join(f"{u} ({unit_names.get(u, u)})" for u in col.options) + ".")
        elif col.type == "enum":
            guide.append("Choose from the list.")
        if col.role == "data_basis":
            guide.append(
                "Sets the data-quality grade: "
                + "; ".join(f"{k} = {v}" for k, v in terminal.quality_grades.items())
                + "."
            )
        rows.append((col.header, " ".join(guide)))
    rows.append((None, None))
    rows.append(
        (
            "Missing data",
            "Leave a period out rather than entering 0. Missing periods are reported as "
            "Missing; a 0 is reported as a real zero.",
        )
    )
    for r, (a, b) in enumerate(rows, start=1):
        ins.cell(row=r, column=1, value=a).font = Font(bold=True)
        ins.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    ins.cell(row=1, column=1).font = Font(bold=True, size=14)

    for i, (k, v) in enumerate(
        [
            ("terminal_key", terminal.key),
            ("template_version", terminal.template_version),
            ("generated_at", datetime.now(UTC).isoformat()),
        ],
        start=1,
    ):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)
    lists.sheet_state = "hidden"
    meta.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
