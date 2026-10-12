"""Bills (PRD F5 Method 2) and AI-assisted categorization (PRD F6) with a deterministic fake provider.

The real provider is never called in tests; the fake stands in for whatever an AI would propose.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.ai.provider import BillExtraction, CategorySuggestion, NullProvider
from tests.conftest import Rest, bearer

pytestmark = pytest.mark.db


class FakeProvider:
    name = "fake"
    model = "fake-model"
    available = True

    def __init__(self, suggest_path_contains: str | None = None):
        self.suggest_path_contains = suggest_path_contains

    def extract_bill(self, data, media_type, fields):
        return BillExtraction(
            "electricity_bill",
            {
                "period_start": "2025-01-01",
                "period_end": "2025-01-31",
                "quantity": "125000",
                "amount_billed": "1050000",
                "meter_id": "M-77",
                "supplier_name": "DESCO",
            },
            {"quantity": "high", "period_start": "high"},
            None,
        )

    def suggest_category(self, terminal_label, inputs, options):
        hit = next((o for o in options if self.suggest_path_contains and self.suggest_path_contains in o["path"]), None)
        return CategorySuggestion(hit["id"], "fake reason") if hit else CategorySuggestion("not-an-option", "x")


@pytest.fixture
def project(engine, tenant_b):
    admin = bearer(tenant_b.tokens["COMPANY_ADMIN"])
    year = 2026
    r = engine.post(
        "/projects",
        headers=admin,
        json={
            "company_id": tenant_b.company_id,
            "name": "AI tests",
            "period_type": "year",
            "boundary": "division",
            "division_name": f"ai-{id(object())}",
            "start_date": f"{year}-01-01",
            "end_date": f"{year}-12-31",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_bill_is_never_saved_before_confirmation(engine, tenant_b, project, monkeypatch):
    monkeypatch.setattr("app.api.bills.get_provider", lambda: FakeProvider())
    h = bearer(tenant_b.tokens["DATA_ENTRY"])
    r = engine.post(
        f"/projects/{project}/terminals/s2_electricity/bills",
        headers=h,
        files={"file": ("bill.pdf", b"%PDF-1.4 fake", "application/pdf")},
    )
    assert r.status_code == 201 and r.json()["status"] == "AI_EXTRACTED"
    fid = r.json()["file_id"]
    rest = Rest(tenant_b.tokens["DATA_ENTRY"])
    assert rest.get("activity_records", file_id=f"eq.{fid}").json() == []
    assert rest.get("v_file_status", file_id=f"eq.{fid}").json()[0]["status"] == "AWAITING_CONFIRMATION"

    bill = engine.get(f"/files/{fid}/bill", headers=h).json()
    assert bill["extraction"]["status"] == "AI_EXTRACTED"
    assert bill["extraction"]["extracted"]["fields"]["quantity"] == "125000"

    # The user corrects the kWh before confirming; the confirmed value is what gets saved.
    r = engine.post(
        f"/files/{fid}/bill/confirm",
        headers=h,
        json={
            "facility_id": tenant_b.facility_id,
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "supplier_name": "DESCO",
            "meter_id": "M-77",
            "quantity": "124500",
            "unit": "kWh",
            "amount_billed": "1050000",
        },
    )
    assert r.status_code == 200, r.text
    rec = rest.get("activity_records", file_id=f"eq.{fid}").json()
    assert len(rec) == 1 and rec[0]["source_type"] == "bill_upload" and rec[0]["review_status"] == "USER_VERIFIED"
    assert float(rec[0]["original_quantity"]) == 124500 and rec[0]["quality_grade"] == "A"
    ex = rest.get("bill_extractions", file_id=f"eq.{fid}").json()[0]
    assert ex["status"] == "USER_VERIFIED" and ex["confirmed_by"] is not None
    assert (
        engine.post(
            f"/files/{fid}/bill/confirm",
            headers=h,
            json={
                "facility_id": tenant_b.facility_id,
                "period_start": "2026-01-01",
                "period_end": "2026-01-31",
                "supplier_name": "DESCO",
                "quantity": "1",
            },
        ).status_code
        == 409
    )


def test_confirmed_bill_values_are_validated(engine, tenant_b, project, monkeypatch):
    monkeypatch.setattr("app.api.bills.get_provider", lambda: FakeProvider())
    h = bearer(tenant_b.tokens["DATA_ENTRY"])
    fid = engine.post(
        f"/projects/{project}/terminals/s2_electricity/bills",
        headers=h,
        files={"file": ("bill.png", b"\x89PNG fake", "image/png")},
    ).json()["file_id"]
    r = engine.post(
        f"/files/{fid}/bill/confirm",
        headers=h,
        json={
            "facility_id": tenant_b.facility_id,
            "period_start": "2025-01-01",
            "period_end": "2025-01-31",
            "supplier_name": "DESCO",
            "quantity": "100",
            "unit": "kWh",
        },
    )
    assert r.status_code == 422 and "outside the project period" in r.json()["detail"]


def test_without_ai_the_bill_form_is_manual(engine, tenant_b, project, monkeypatch):
    monkeypatch.setattr("app.api.bills.get_provider", lambda: NullProvider())
    h = bearer(tenant_b.tokens["DATA_ENTRY"])
    r = engine.post(
        f"/projects/{project}/terminals/s2_electricity/bills",
        headers=h,
        files={"file": ("bill.pdf", b"%PDF-1.4", "application/pdf")},
    )
    assert r.json()["status"] == "FAILED"
    bill = engine.get(f"/files/{r.json()['file_id']}/bill", headers=h).json()
    assert all(v is None for v in bill["extraction"]["extracted"]["fields"].values())


def test_bills_only_on_terminals_that_accept_them(engine, tenant_b, project):
    r = engine.post(
        f"/projects/{project}/terminals/s1_mobile/bills",
        headers=bearer(tenant_b.tokens["DATA_ENTRY"]),
        files={"file": ("bill.pdf", b"%PDF", "application/pdf")},
    )
    assert r.status_code == 409


def _upload_fugitive(engine, token, project, gas):
    tpl = engine.get(f"/projects/{project}/terminals/s1_fugitive/template", headers=bearer(token)).content
    wb = load_workbook(io.BytesIO(tpl))
    ws = wb["Data"]
    for c, v in enumerate(
        ["Plant 1", date(2026, 1, 1), date(2026, 1, 31), "CoolServ", "Chiller", gas, 5, "kg", "Service records"],
        start=1,
    ):
        ws.cell(row=2, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return engine.post(
        f"/projects/{project}/terminals/s1_fugitive/uploads",
        headers=bearer(token),
        files={"file": ("gas.xlsx", buf.getvalue(), "application/octet-stream")},
    )


def test_ai_suggests_only_where_rules_find_nothing_and_only_within_the_terminal(engine, tenant_b, project, monkeypatch):
    token = tenant_b.tokens["DATA_ENTRY"]
    monkeypatch.setattr("app.categorize.service.get_provider", lambda: FakeProvider("R-410A"))
    assert _upload_fugitive(engine, token, project, "HFC-134a").json()["records_created"] == 1
    assert _upload_fugitive(engine, token, project, "Other gas").json()["records_created"] == 1
    cats = {
        c["input_display"]["Refrigerant / gas"]: c
        for c in engine.get(f"/projects/{project}/terminals/s1_fugitive/categorizations", headers=bearer(token)).json()
    }
    assert cats["HFC-134a"]["suggested_by"] == "rule"
    assert cats["Other gas"]["suggested_by"] == "ai" and cats["Other gas"]["suggested_path"].endswith("R-410A")
    assert cats["Other gas"]["status"] == "SUGGESTED", "AI suggestions still need a person to accept them"


def test_ai_answers_outside_the_option_list_are_discarded(engine, tenant_b, project, monkeypatch):
    token = tenant_b.tokens["DATA_ENTRY"]
    monkeypatch.setattr("app.categorize.service.get_provider", lambda: FakeProvider(None))
    _upload_fugitive(engine, token, project, "Other gas")
    cat = engine.get(f"/projects/{project}/terminals/s1_fugitive/categorizations", headers=bearer(token)).json()[0]
    assert cat["status"] == "NEEDS" and cat["suggested_by"] is None
