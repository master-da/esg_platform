"""Templates, parsing, terminal detection and validation (PRD F5, F5.2). Pure unit tests.

Every terminal configuration in reference-data/terminals is exercised, so adding a terminal adds
coverage automatically.
"""

from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.categorize.rules import input_key, match_rule
from app.ingest.parser import ParseError, detect_terminal, parse_file
from app.ingest.templates import DATA_SHEET, build_template
from app.ingest.terminal import Terminal
from app.ingest.validation import ValidationContext, validate

TERMINAL_DIR = Path(__file__).resolve().parents[2] / "reference-data" / "terminals"
TERMINALS = {
    t.key: t for t in (Terminal.from_row(json.loads(p.read_text())) for p in sorted(TERMINAL_DIR.glob("*.json")))
}
UNITS = {
    "kg": "mass",
    "t": "mass",
    "L": "volume",
    "m3": "volume",
    "gal_us": "volume",
    "ML": "volume",
    "kWh": "energy",
    "MWh": "energy",
    "GJ": "energy",
    "MMBtu": "energy",
    "km": "distance",
    "mile": "distance",
    "pkm": "passenger_distance",
    "tkm": "freight_distance",
    "BDT": "currency",
    "USD": "currency",
    "room_night": "count",
    "unit": "count",
    "hour": "time",
    "day": "time",
}
ALIASES = {**{u.casefold(): u for u in UNITS}, "litres": "L", "tonnes": "t"}
CTX = ValidationContext(
    facilities={
        "gazipur factory": ("fac-1", "Gazipur Factory"),
        "company-level (all sites)": ("fac-0", "Company-level (all sites)"),
    },
    unit_dimensions=UNITS,
    unit_aliases=ALIASES,
    project_start=date(2025, 1, 1),
    project_end=date(2025, 12, 31),
)


def sample_value(col, terminal: Terminal):
    if col.type == "facility":
        return "Gazipur Factory"
    if col.type == "date":
        return date(2025, 1, 1) if col.key.endswith("start") else date(2025, 1, 31)
    if col.type in ("decimal", "integer"):
        return 1200
    if col.type == "unit":
        # Respect the terminal's own unit rules for the sample category value.
        for rule in terminal.rules:
            if rule["type"] == "unit_by_value":
                chosen = terminal.column(rule["column"]).options[0]
                return rule["allowed"][chosen][0]
        return col.options[0]
    if col.type == "enum":
        return col.options[0]
    return "Sample text"


def filled_template(terminal: Terminal, overrides: dict | None = None, rows: int = 2) -> bytes:
    wb = load_workbook(io.BytesIO(build_template(terminal, ["Gazipur Factory"], {"L": "litres"})))
    ws = wb[DATA_SHEET]
    for r in range(2, 2 + rows):
        for i, col in enumerate(terminal.columns, start=1):
            value = (overrides or {}).get(col.key, sample_value(col, terminal))
            ws.cell(row=r, column=i, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("key", sorted(TERMINALS))
def test_template_round_trip_for_every_terminal(key):
    t = TERMINALS[key]
    parsed = parse_file("upload.xlsx", filled_template(t), 1000)
    assert detect_terminal(parsed, TERMINALS) == key
    rows, issues = validate(parsed, t, CTX)
    assert issues == []
    assert len(rows) == 2 and rows[0].source_row == 2
    q = t.role("quantity")
    assert rows[0].values[q.key] == Decimal("1200")


@pytest.mark.parametrize("key", sorted(TERMINALS))
def test_template_has_dropdowns_and_hidden_meta(key):
    wb = load_workbook(io.BytesIO(build_template(TERMINALS[key], ["Gazipur Factory"], {})))
    assert wb["_meta"].sheet_state == "hidden" and wb["_meta"]["B1"].value == key
    assert wb[DATA_SHEET].data_validations.dataValidation, "template must carry dropdown validation"


def test_wrong_terminal_is_detected_and_named():
    mobile = filled_template(TERMINALS["s1_mobile"])
    parsed = parse_file("S1_Mobile.xlsx", mobile, 1000)
    detected = detect_terminal(parsed, TERMINALS)
    assert detected == "s1_mobile" and detected != "s2_electricity"


def test_csv_without_meta_is_detected_by_headers():
    t = TERMINALS["s2_electricity"]
    header = ",".join(c.header for c in t.columns)
    row = ",".join(
        [
            "Gazipur Factory",
            "2025-01-01",
            "2025-01-31",
            "DESCO",
            "M-1",
            "Grid electricity",
            "125000",
            "kWh",
            "1000000",
            "Utility bill / meter reading",
            "",
        ]
    )
    parsed = parse_file("bill.csv", f"{header}\n{row}\n".encode(), 1000)
    assert detect_terminal(parsed, TERMINALS) == "s2_electricity"
    rows, issues = validate(parsed, t, CTX)
    assert issues == [] and rows[0].values["quantity"] == Decimal("125000")


def test_validation_issues_name_row_and_field():
    t = TERMINALS["s1_mobile"]
    bad = filled_template(t, {"fuel_type": "Petrol", "unit": "km", "quantity": None}, rows=1)
    rows, issues = validate(parse_file("x.xlsx", bad, 1000), t, CTX)
    by_field = {i.field: i for i in issues}
    assert by_field["Fuel type"].row == 2 and "not in the fuel type list" in by_field["Fuel type"].problem
    assert by_field["Fuel quantity"].problem == "Required field is empty"
    assert by_field["Unit"].row == 2


def test_unit_dimension_rule_rejects_distance_for_fuel_terminal():
    t = TERMINALS["s1_mobile"]
    column = next(c for c in t.columns if c.key == "unit")
    assert "km" not in column.options  # dropdown already excludes it; the rule guards hand-edited files
    from dataclasses import replace

    loose = replace(
        t, columns=tuple(replace(c, options=(*c.options, "km")) if c.key == "unit" else c for c in t.columns)
    )
    _, issues = validate(parse_file("x.xlsx", filled_template(loose, {"unit": "km"}, rows=1), 1000), loose, CTX)
    assert any("non-fuel" in i.problem.lower() or "distance" in i.problem.lower() for i in issues)


def test_period_outside_project_is_rejected():
    t = TERMINALS["s2_electricity"]
    data = filled_template(t, {"period_start": date(2024, 12, 1), "period_end": date(2024, 12, 31)}, rows=1)
    _, issues = validate(parse_file("x.xlsx", data, 1000), t, CTX)
    assert any("outside the project period" in i.problem for i in issues)


def test_unknown_facility_and_bad_number():
    t = TERMINALS["s2_electricity"]
    data = filled_template(t, {"facility": "Chittagong Plant", "quantity": "twelve"}, rows=1)
    _, issues = validate(parse_file("x.xlsx", data, 1000), t, CTX)
    problems = " ".join(i.problem for i in issues)
    assert 'Facility "Chittagong Plant" does not exist' in problems and '"twelve" is not a number' in problems


def test_missing_required_column_rejects_the_file():
    t = TERMINALS["s2_electricity"]
    header = ",".join(c.header for c in t.columns if c.key != "quantity")
    _, issues = validate(parse_file("x.csv", f"{header}\n".encode(), 1000), t, CTX)
    assert issues and issues[0].row is None and "missing" in issues[0].problem


def test_dd_mm_yyyy_dates_and_thousands_separators_are_accepted():
    t = TERMINALS["s2_electricity"]
    header = ",".join(c.header for c in t.columns)
    row = [
        "Gazipur Factory",
        "01/03/2025",
        "31/03/2025",
        "DPDC",
        "",
        "Grid electricity",
        '"127,000"',
        "kWh",
        "",
        "Estimate",
        "",
    ]
    rows, issues = validate(parse_file("x.csv", f"{header}\n{','.join(row)}\n".encode(), 1000), t, CTX)
    assert issues == []
    assert rows[0].values["period_start"] == date(2025, 3, 1) and rows[0].values["quantity"] == Decimal("127000")


def test_unsupported_file_type():
    with pytest.raises(ParseError):
        parse_file("data.pdf", b"%PDF", 1000)


def test_categorization_rules_match_normalised_input():
    t = TERMINALS["s1_mobile"]
    key = input_key(t, {"vehicle_type": "Pick-up truck", "fuel_type": "Diesel"})
    assert key == "vehicle_type=pick-up truck|fuel_type=diesel"
    rules = [
        {"match": {"fuel_type": "diesel"}, "priority": 100, "taxonomy_code": "A"},
        {"match": {"fuel_type": "diesel", "vehicle_type": "pick-up truck"}, "priority": 100, "taxonomy_code": "B"},
    ]
    assert match_rule(rules, key)["taxonomy_code"] == "B"
    assert match_rule(rules, input_key(t, {"vehicle_type": "Car", "fuel_type": "CNG"})) is None
