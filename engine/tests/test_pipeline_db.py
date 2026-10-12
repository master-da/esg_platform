"""End-to-end data pipeline against the local stack (PRD F3, F4, F5, F5.1, F5.2, F6, F7).

Uses the seeded, cited DESNZ 2025 diesel factor for mobile combustion (2.66155 kg CO2e/L).
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests.conftest import Rest, bearer

pytestmark = pytest.mark.db


def fill(template: bytes, rows: list[list]) -> bytes:
    wb = load_workbook(io.BytesIO(template))
    ws = wb["Data"]
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(engine, token, project_id, terminal, name, data):
    return engine.post(
        f"/projects/{project_id}/terminals/{terminal}/uploads",
        headers=bearer(token),
        files={"file": (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


@pytest.fixture(scope="module")
def project(engine, tenant_a):
    r = engine.post(
        "/projects",
        headers=bearer(tenant_a.tokens["SUSTAINABILITY_MANAGER"]),
        json={
            "company_id": tenant_a.company_id,
            "name": "FY 2025",
            "period_type": "year",
            "start_date": "2025-01-01",
            "end_date": "2025-12-31",
            "upload_deadline": "2026-03-31",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_projects_cannot_overlap(engine, tenant_a, project):
    r = engine.post(
        "/projects",
        headers=bearer(tenant_a.tokens["SUSTAINABILITY_MANAGER"]),
        json={
            "company_id": tenant_a.company_id,
            "name": "Q3 2025",
            "period_type": "quarter",
            "start_date": "2025-07-01",
            "end_date": "2025-09-30",
        },
    )
    assert r.status_code == 409 and "overlaps" in r.json()["detail"]


def test_read_only_roles_cannot_create_projects_or_upload(engine, tenant_a, project):
    r = engine.post(
        "/projects",
        headers=bearer(tenant_a.tokens["VIEWER"]),
        json={
            "company_id": tenant_a.company_id,
            "name": "Viewer project",
            "period_type": "year",
            "start_date": "2030-01-01",
            "end_date": "2030-12-31",
        },
    )
    assert r.status_code == 403
    tpl = engine.get(f"/projects/{project}/terminals/s1_mobile/template", headers=bearer(tenant_a.tokens["VIEWER"]))
    assert tpl.status_code == 200
    r = upload(engine, tenant_a.tokens["AUDITOR"], project, "s1_mobile", "x.xlsx", tpl.content)
    assert r.status_code == 403, r.text


def test_other_tenant_cannot_see_or_touch_the_project(engine, tenant_b, project):
    h = bearer(tenant_b.tokens["COMPANY_ADMIN"])
    assert engine.get(f"/projects/{project}", headers=h).status_code == 404
    assert engine.get(f"/projects/{project}/results", headers=h).status_code == 404
    assert Rest(tenant_b.tokens["COMPANY_ADMIN"]).get("projects", id=f"eq.{project}").json() == []


def test_full_pipeline(engine, tenant_a, internal_user, project, superuser_conn):
    de = tenant_a.tokens["DATA_ENTRY"]
    h = bearer(de)
    tpl = engine.get(f"/projects/{project}/terminals/s1_mobile/template", headers=h).content

    # 1. A file with errors is rejected with per-row, per-field messages; no records are created.
    bad = fill(
        tpl,
        [
            [
                "Plant 1",
                date(2025, 1, 1),
                date(2025, 1, 31),
                "Padma Oil",
                "Pick-up truck",
                "Diesel",
                100,
                "L",
                "Fuel receipts / fuel card",
            ],
            ["Plant 1", date(2025, 2, 1), date(2025, 2, 28), "Padma Oil", "Car", "Petrol", None, "L", "Estimate"],
        ],
    )
    r = upload(engine, de, project, "s1_mobile", "S1_Mobile_Feb.xlsx", bad)
    assert r.status_code == 201 and r.json()["validation_passed"] is False and r.json()["records_created"] == 0
    issues = engine.get(f"/files/{r.json()['file_id']}/issues", headers=h).json()
    assert {(i["row_number"], i["field"]) for i in issues} == {(3, "Fuel type"), (3, "Fuel quantity")}

    # 2. A template from another terminal is rejected naming the right page.
    elec_tpl = engine.get(f"/projects/{project}/terminals/s2_electricity/template", headers=h).content
    r = upload(engine, de, project, "s1_mobile", "electricity.xlsx", elec_tpl)
    assert r.status_code == 422 and "Purchased electricity" in r.json()["detail"]

    # 3. A good file creates one pending record per row, and rules suggest a category.
    good = fill(
        tpl,
        [
            [
                "Plant 1",
                date(2025, 1, 1),
                date(2025, 1, 31),
                "Padma Oil",
                "Pick-up truck",
                "Diesel",
                1200,
                "L",
                "Fuel receipts / fuel card",
            ],
            [
                "Plant 1",
                date(2025, 2, 1),
                date(2025, 2, 28),
                "Padma Oil",
                "Pick-up truck",
                "Diesel",
                1.5,
                "t",
                "Supplier statement",
            ],
        ],
    )
    r = upload(engine, de, project, "s1_mobile", "S1_Mobile_Jan.xlsx", good)
    assert r.status_code == 201 and r.json()["records_created"] == 2
    file_id = r.json()["file_id"]
    recs = engine.get(f"/projects/{project}/terminals/s1_mobile/records", headers=h).json()
    assert {x["review_status"] for x in recs} == {"pending_review"}
    assert any(x["normalized_unit"] == "kg" and Decimal(str(x["normalized_quantity"])) == 1500 for x in recs)
    cats = engine.get(f"/projects/{project}/terminals/s1_mobile/categorizations", headers=h).json()
    assert len(cats) == 1 and cats[0]["status"] == "SUGGESTED" and cats[0]["suggested_by"] == "rule"
    assert cats[0]["suggested_path"] == "Scope 1 > Mobile combustion > Road > Diesel"
    assert "code" not in str(cats[0].keys())
    status = Rest(de).get("v_file_status", file_id=f"eq.{file_id}").json()[0]["status"]
    assert status == "READY_TO_CATEGORIZE", "nothing is calculated before a person accepts"

    # 4. Accepting calculates deterministically and marks rows user-verified.
    r = engine.post(f"/projects/{project}/terminals/s1_mobile/categorizations/accept-suggested", headers=h)
    assert r.json() == {"accepted": 1, "calculated": 2, "unknown_factor": 0, "unit_error": 0}
    results = engine.get(f"/projects/{project}/results", headers=h).json()
    mobile = next(c for c in results if c["terminal_key"] == "s1_mobile")
    jan = next(x for x in mobile["rows"] if x["input_unit"] == "L")
    assert Decimal(str(jan["co2e_t"])) == Decimal("3.19386")  # 1200 L x 2.66155 kg/L
    feb = next(x for x in mobile["rows"] if x["input_unit"] == "t")
    assert Decimal(str(feb["co2e_t"])) == Decimal("4.805867145")  # 1.5 t x 3203.91143 kg/t
    trail = engine.get(f"/calculations/{feb['calculation_id']}", headers=h).json()
    assert trail["file_name"] == "S1_Mobile_Jan.xlsx" and trail["source_row"] == 3
    assert trail["conversion_from"] == "t" and trail["conversion_to"] == "L"
    assert trail["factor_source_name"].startswith("UK DESNZ") and trail["factor_source_year"] == 2025
    assert "factor_code" not in trail, "factor codes are internal only"
    assert (
        "01020000010001"
        in engine.get(f"/calculations/{feb['calculation_id']}", headers=bearer(internal_user)).json()["factor_code"]
    )
    recs = engine.get(f"/projects/{project}/terminals/s1_mobile/records", headers=h).json()
    assert {x["review_status"] for x in recs} == {"USER_VERIFIED"}
    assert Rest(de).get("v_file_status", file_id=f"eq.{file_id}").json()[0]["status"] == "PRELIMINARY"

    # 5. Only internal users can confirm "Calculated" or download the calculated file.
    assert engine.post(f"/files/{file_id}/review", headers=h, json={"action": "confirm"}).status_code == 403
    assert engine.get(f"/files/{file_id}/calculated", headers=h).status_code == 403
    assert engine.get(f"/files/{file_id}/calculated", headers=bearer(internal_user)).status_code == 200
    assert (
        engine.post(f"/files/{file_id}/review", headers=bearer(internal_user), json={"action": "confirm"}).status_code
        == 200
    )
    assert Rest(de).get("v_file_status", file_id=f"eq.{file_id}").json()[0]["status"] == "CALCULATED"

    # The rejected file still blocks approval; delete it (audit keeps the history).
    bad_id = next(
        f["file_id"]
        for f in Rest(de).get("v_file_status", project_id=f"eq.{project}").json()
        if f["original_name"] == "S1_Mobile_Feb.xlsx"
    )
    assert engine.delete(f"/files/{bad_id}", headers=h).status_code == 200
    deleted = superuser_conn.execute(
        "select count(*) as n from public.audit_log where entity_id = %s and action = 'DELETE_FILE'", (bad_id,)
    ).fetchone()
    assert deleted["n"] == 1

    # 6. Stage workflow: DATA_ENTRY cannot approve; approvers can approve and lock.
    sm = bearer(tenant_a.tokens["SUSTAINABILITY_MANAGER"])
    for to in ("DATA_COLLECTION", "REVIEW"):
        assert engine.post(f"/projects/{project}/transition", headers=h, json={"to": to}).status_code == 200
    assert engine.post(f"/projects/{project}/transition", headers=h, json={"to": "APPROVED"}).status_code == 403
    assert engine.post(f"/projects/{project}/transition", headers=sm, json={"to": "APPROVED"}).status_code == 200
    assert engine.post(f"/projects/{project}/transition", headers=sm, json={"to": "LOCKED"}).status_code == 200

    # 7. Locked data cannot be edited; the message is explicit.
    r = upload(engine, de, project, "s1_mobile", "late.xlsx", good)
    assert r.status_code >= 400
    with pytest.raises(Exception, match="LOCKED"):
        superuser_conn.execute(
            "update public.activity_records set original_quantity = 1 where file_id = %s", (file_id,)
        )

    # 8. Recalculating a locked project needs explicit confirmation and creates a new version.
    r = engine.post(f"/projects/{project}/recalculate", headers=sm, json={"reason": "Factor review"})
    assert r.status_code == 409
    r = engine.post(f"/projects/{project}/recalculate", headers=sm, json={"reason": "Factor review", "confirm": True})
    assert r.status_code == 200 and r.json()["version_no"] == 2, r.text
    detail = engine.get(f"/projects/{project}", headers=h).json()
    assert [v["version_no"] for v in detail["versions"]] == [2, 1]
    assert Decimal(str(detail["co2e_t"])).quantize(Decimal("0.00001")) == Decimal("7.99973")
    old = superuser_conn.execute(
        """select count(*) as n from public.emission_calculations c join public.calculation_versions v
           on v.id = c.calc_version_id where v.project_id = %s and v.version_no = 1""",
        (project,),
    ).fetchone()
    assert old["n"] == 2, "version 1 results are preserved"


def test_unknown_factor_is_never_zero(engine, tenant_b):
    """Bangladesh grid factors exist only up to June 2022: a 2025 electricity row stays UNKNOWN."""
    admin = bearer(tenant_b.tokens["COMPANY_ADMIN"])
    pid = engine.post(
        "/projects",
        headers=admin,
        json={
            "company_id": tenant_b.company_id,
            "name": "FY 2025 B",
            "period_type": "year",
            "start_date": "2025-01-01",
            "end_date": "2025-12-31",
        },
    ).json()["id"]
    tpl = engine.get(f"/projects/{pid}/terminals/s2_electricity/template", headers=admin).content
    data = fill(
        tpl,
        [
            [
                "Plant 1",
                date(2025, 1, 1),
                date(2025, 1, 31),
                "DESCO",
                "M-1",
                "Grid electricity",
                125000,
                "kWh",
                None,
                "Utility bill / meter reading",
            ]
        ],
    )
    r = upload(engine, tenant_b.tokens["COMPANY_ADMIN"], pid, "s2_electricity", "Jan.xlsx", data)
    assert r.json()["records_created"] == 1
    cats = engine.get(f"/projects/{pid}/terminals/s2_electricity/categorizations", headers=admin).json()
    assert cats[0]["factor_availability"] == "none"
    res = engine.post(f"/categorizations/{cats[0]['id']}/accept", headers=admin, json={}).json()
    assert res["unknown_factor"] == 1 and res["calculated"] == 0
    rows = engine.get(f"/projects/{pid}/results", headers=admin).json()[0]["rows"]
    assert rows[0]["status"] == "UNKNOWN_FACTOR" and rows[0]["co2e_t"] is None
    status = Rest(tenant_b.tokens["COMPANY_ADMIN"]).get("v_file_status", project_id=f"eq.{pid}").json()[0]["status"]
    assert status == "NEED_ATTENTION"


def test_missing_months_are_reported_not_zero(engine, tenant_a):
    """A facility that reported January and March but not February shows February as Missing."""
    h = bearer(tenant_a.tokens["SUSTAINABILITY_MANAGER"])
    pid = engine.post(
        "/projects",
        headers=h,
        json={
            "company_id": tenant_a.company_id,
            "name": "Q1 coverage",
            "period_type": "quarter",
            "boundary": "division",
            "division_name": "coverage",
            "start_date": "2025-01-01",
            "end_date": "2025-03-31",
        },
    ).json()["id"]
    tpl = engine.get(f"/projects/{pid}/terminals/s1_mobile/template", headers=h).content
    rows = [
        ["Plant 1", date(2025, m, 1), date(2025, m, 28), "Padma", "Car", "Diesel", 0 if m == 3 else 10, "L", "Estimate"]
        for m in (1, 3)
    ]
    assert (
        upload(engine, tenant_a.tokens["DATA_ENTRY"], pid, "s1_mobile", "q1.xlsx", fill(tpl, rows)).status_code == 201
    )
    cov = engine.get(f"/projects/{pid}/coverage", headers=h).json()
    plant = next(c for c in cov if c["terminal_key"] == "s1_mobile")
    assert plant["missing_months"] == ["2025-02"], "an explicit 0 in March is data; February is Missing"
