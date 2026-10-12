"""Compliance, reporting and dashboards (PRD Features 14, 15, 16)."""

from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date, timedelta
from decimal import ROUND_HALF_EVEN, Decimal

import psycopg
import pytest
from openpyxl import load_workbook

from tests.conftest import Rest, bearer

pytestmark = pytest.mark.db


@pytest.fixture(scope="module")
def locked_project(engine, tenant_a, internal_user):
    """A small, fully calculated and locked project with data in two facilities."""
    sm = bearer(tenant_a.tokens["SUSTAINABILITY_MANAGER"])
    pid = engine.post(
        "/projects",
        headers=sm,
        json={
            "company_id": tenant_a.company_id,
            "name": "Reporting",
            "period_type": "year",
            "boundary": "division",
            "division_name": f"rep-{uuid.uuid4().hex[:6]}",
            "start_date": "2025-01-01",
            "end_date": "2025-12-31",
        },
    ).json()["id"]
    for fac, qty, unit in (("Plant 1", 1200, "L"), ("Company-level (all sites)", 1.5, "t"), ("Plant 1", 300, "L")):
        r = engine.post(
            f"/projects/{pid}/terminals/s1_mobile/manual",
            headers=sm,
            json={
                "values": {
                    "facility": fac,
                    "period_start": "2025-03-01",
                    "period_end": "2025-03-31",
                    "supplier_name": "Padma Oil",
                    "vehicle_type": "Truck",
                    "fuel_type": "Diesel",
                    "quantity": qty,
                    "unit": unit,
                    "data_basis": "Fuel receipts / fuel card",
                }
            },
        )
        assert r.status_code == 201, r.text
    engine.post(f"/projects/{pid}/terminals/s1_mobile/categorizations/accept-suggested", headers=sm)
    for f in Rest(tenant_a.tokens["SUSTAINABILITY_MANAGER"]).get("v_file_status", project_id=f"eq.{pid}").json():
        assert (
            engine.post(
                f"/files/{f['file_id']}/review", headers=bearer(internal_user), json={"action": "confirm"}
            ).status_code
            == 200
        )
    rest = Rest(tenant_a.tokens["SUSTAINABILITY_MANAGER"])
    assert (
        rest.post(
            "assumptions",
            {
                "company_id": tenant_a.company_id,
                "project_id": pid,
                "category": "estimate",
                "statement": "Company-level fuel is allocated from the annual supplier statement.",
            },
        ).status_code
        == 201
    )
    for to in ("DATA_COLLECTION", "REVIEW", "APPROVED", "LOCKED"):
        assert engine.post(f"/projects/{pid}/transition", headers=sm, json={"to": to}).status_code == 200, to
    return pid


def export(engine, token, **body):
    return engine.post("/exports", headers=bearer(token), json=body)


# -- F14 C: audit export -------------------------------------------------------------------


def test_auditor_can_recompute_every_result_from_the_audit_export_alone(engine, tenant_a, locked_project):
    r = export(engine, tenant_a.tokens["AUDITOR"], project_id=locked_project, export_type="audit", format="xlsx")
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Calculations"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, (c.value for c in row), strict=False)) for row in ws.iter_rows(min_row=2)]
    assert len(rows) == 3
    for row in rows:
        qty = Decimal(str(row["Input value"]))
        conv = Decimal(str(row["Conversion factor"])) if row["Conversion factor"] else Decimal(1)
        factor = Decimal(str(row["Emission factor (kg CO2e per unit)"]))
        recomputed = (qty * conv * factor).quantize(Decimal("0.000001"), rounding=ROUND_HALF_EVEN)
        assert recomputed == Decimal(str(row["Calculated GHG (kg CO2e)"])), row
        assert row["Category (description path)"].startswith("Scope 1 > Mobile combustion")
        assert row["Emission factor database"].startswith("UK DESNZ") and row["Database year"] == 2025
    text = " ".join(str(c.value) for s in wb.worksheets for row in s.iter_rows() for c in row if c.value is not None)
    assert not re.search(r"\b0[1-3]\d{12}\b", text), "factor codes must not appear in the audit export"
    assert "factor_code" not in text
    for sheet in ("Version history", "Factor sources", "Assumptions & quality"):
        assert sheet in wb.sheetnames


def test_audit_csv_has_the_same_rows(engine, tenant_a, locked_project):
    r = export(engine, tenant_a.tokens["AUDITOR"], project_id=locked_project, export_type="audit", format="csv")
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert len(rows) == 3 and "Calculated GHG (t CO2e)" in rows[0]


# -- F14 C: PDF ----------------------------------------------------------------------------


def test_pdf_report_for_locked_project_with_assumptions(engine, tenant_a, locked_project):
    r = export(engine, tenant_a.tokens["VIEWER"], project_id=locked_project, export_type="GHG_PROTOCOL", format="pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    from pypdf import PdfReader

    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(r.content)).pages)
    assert "Assumptions and limitations" in text
    assert "allocated from the annual supplier" in text, "assumptions come from the Assumptions Register"
    assert not re.search(r"\b(certified|audited|verified)\b", text, re.IGNORECASE)


def test_pdf_wording_guard_blocks_verification_claims(engine, tenant_a, internal_user, locked_project, superuser_conn):
    superuser_conn.execute("set session_replication_role = replica")  # the project is locked; add a claim directly
    superuser_conn.execute(
        "insert into public.assumptions (company_id, project_id, category, statement) values (%s, %s, 'other', %s)",
        (tenant_a.company_id, locked_project, "Totals were verified by the finance team."),
    )
    superuser_conn.execute("set session_replication_role = origin")
    r = export(engine, tenant_a.tokens["VIEWER"], project_id=locked_project, export_type="GRI", format="pdf")
    assert r.status_code == 422 and "verified" in r.json()["detail"]
    # A recorded third-party verification (internal users only) lifts the restriction.
    assert Rest(tenant_a.tokens["COMPANY_ADMIN"]).post(
        "verifications",
        {
            "company_id": tenant_a.company_id,
            "project_id": locked_project,
            "verifier": "X",
            "standard": "ISO 14064-3",
            "assurance_level": "limited",
            "statement_date": "2026-03-01",
            "scope": "Scope 1",
        },
    ).status_code in (401, 403)
    assert (
        Rest(internal_user)
        .post(
            "verifications",
            {
                "company_id": tenant_a.company_id,
                "project_id": locked_project,
                "verifier": "Assurance Co",
                "standard": "ISO 14064-3",
                "assurance_level": "limited",
                "statement_date": "2026-03-01",
                "scope": "Scope 1",
            },
        )
        .status_code
        == 201
    )
    assert (
        export(
            engine, tenant_a.tokens["VIEWER"], project_id=locked_project, export_type="GRI", format="pdf"
        ).status_code
        == 200
    )


def test_pdf_only_for_approved_or_locked(engine, tenant_b):
    h = bearer(tenant_b.tokens["COMPANY_ADMIN"])
    pid = engine.post(
        "/projects",
        headers=h,
        json={
            "company_id": tenant_b.company_id,
            "name": "Draft",
            "period_type": "year",
            "boundary": "division",
            "division_name": f"d-{uuid.uuid4().hex[:5]}",
            "start_date": "2025-01-01",
            "end_date": "2025-12-31",
        },
    ).json()["id"]
    assert (
        export(engine, tenant_b.tokens["COMPANY_ADMIN"], project_id=pid, export_type="GRI", format="pdf").status_code
        == 409
    )


# -- Warnings before export ----------------------------------------------------------------


def test_missing_data_is_warned_before_export(engine, tenant_a, locked_project):
    r = engine.post(
        "/exports/preview",
        headers=bearer(tenant_a.tokens["VIEWER"]),
        json={"project_id": locked_project, "export_type": "GRI", "format": "xlsx"},
    )
    texts = " ".join(w["text"] for w in r.json()["warnings"])
    assert "Missing months" in texts and "no data" in texts


# -- F14 B: a new framework is mapping rows only ---------------------------------------------


def test_new_framework_needs_only_mapping_rows(engine, platform_admin, tenant_a, locked_project):
    h = bearer(platform_admin)
    key = f"TEST_{uuid.uuid4().hex[:6].upper()}"
    assert (
        engine.post("/reference/frameworks", headers=h, json={"key": key, "name": "Test framework"}).status_code == 201
    )
    assert (
        engine.post(
            "/reference/mappings",
            headers=h,
            json={"framework_key": key, "metric_key": "GHG_SCOPE_1", "disclosure_ref": "T-1"},
        ).status_code
        == 201
    )
    r = engine.post(
        "/exports",
        headers=bearer(tenant_a.tokens["VIEWER"]),
        json={"project_id": locked_project, "export_type": key, "format": "csv"},
    )
    assert r.status_code == 200, r.text
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["Disclosure"] for x in rows] == ["T-1"] and rows[0]["Metric"] == "Scope 1 GHG emissions"
    assert (
        engine.post(
            "/exports",
            headers=bearer(tenant_a.tokens["VIEWER"]),
            json={"project_id": locked_project, "export_type": "NOPE", "format": "csv"},
        ).status_code
        == 422
    )


# -- F14 A: requirements -------------------------------------------------------------------


def test_requirements_apply_by_data_rules_and_always_show_source(engine, tenant_a, platform_admin, superuser_conn):
    rest = Rest(tenant_a.tokens["COMPANY_ADMIN"])
    rest.patch(
        "facilities",
        {"environmental_classification": "RED", "employee_count": 120, "facility_type": "MANUFACTURING"},
        id=f"eq.{tenant_a.facility_id}",
    )
    reqs = engine.get(
        f"/companies/{tenant_a.company_id}/requirements", headers=bearer(tenant_a.tokens["VIEWER"])
    ).json()
    applying = {r["code"]: r for r in reqs if r["applies"]}
    assert {"BD-DOE-ECR2023-LOCATION-CLEARANCE", "BD-DOE-ECR2023-ECC", "BD-DIFE-SAFETY-COMMITTEE"} <= set(applying)
    for r in reqs:
        assert r["regulator"] and r["source_url"].startswith("http") and r["effective_date"]
    assert applying["BD-DIFE-SAFETY-COMMITTEE"]["applies_to"][0]["facility"] == "Plant 1"
    # A requirement cannot exist without its authority, source or effective date.
    with pytest.raises(psycopg.errors.NotNullViolation):
        superuser_conn.execute(
            "insert into ref.regulatory_requirements (code, regulator, law, title, effective_date) "
            "values ('TEST-X', 'r', 'l', 't', '2025-01-01')"
        )
    # Changing an applicability rule is a data change: a new version, no deployment.
    r = engine.post(
        "/reference/requirements",
        headers=bearer(platform_admin),
        json={
            "code": "TEST-THRESHOLD",
            "regulator": "Test regulator",
            "law": "Test law",
            "title": "Test requirement",
            "source_url": "https://example.invalid/law",
            "effective_date": "2020-01-01",
            "applicability_conditions": {"all": [{"field": "facility.employee_count", "op": "gte", "value": 1000}]},
        },
    )
    assert r.status_code == 201
    reqs = engine.get(
        f"/companies/{tenant_a.company_id}/requirements", headers=bearer(tenant_a.tokens["VIEWER"])
    ).json()
    assert not next(x for x in reqs if x["code"] == "TEST-THRESHOLD")["applies"]
    engine.post(
        "/reference/requirements",
        headers=bearer(platform_admin),
        json={
            "code": "TEST-THRESHOLD",
            "regulator": "Test regulator",
            "law": "Test law",
            "title": "Test requirement",
            "source_url": "https://example.invalid/law",
            "effective_date": "2020-01-01",
            "applicability_conditions": {"all": [{"field": "facility.employee_count", "op": "gte", "value": 100}]},
        },
    )
    reqs = engine.get(
        f"/companies/{tenant_a.company_id}/requirements", headers=bearer(tenant_a.tokens["VIEWER"])
    ).json()
    latest = next(x for x in reqs if x["code"] == "TEST-THRESHOLD")
    assert latest["version"] == 2 and latest["applies"]


def test_universal_requirement_must_be_explicit(engine, platform_admin):
    r = engine.post(
        "/reference/requirements",
        headers=bearer(platform_admin),
        json={
            "code": "TEST-EMPTY",
            "regulator": "R",
            "law": "L",
            "title": "T",
            "source_url": "https://example.invalid",
            "effective_date": "2020-01-01",
            "applicability_conditions": {"all": []},
        },
    )
    assert r.status_code == 422


# -- F15 permits -----------------------------------------------------------------------------


def test_permit_expiring_within_lead_time_raises_an_alert(engine, tenant_a):
    rest = Rest(tenant_a.tokens["FACILITY_MANAGER"])
    today = date.today()
    for number, expiry, lead in (
        ("P-DUE", today + timedelta(days=30), 60),
        ("P-LATER", today + timedelta(days=200), 60),
        ("P-EXPIRED", today - timedelta(days=3), 30),
    ):
        assert (
            rest.post(
                "permits",
                {
                    "company_id": tenant_a.company_id,
                    "facility_id": tenant_a.facility_id,
                    "permit_type": "Environmental Clearance Certificate",
                    "permit_number": number,
                    "issuing_authority": "Department of Environment",
                    "expiry_date": str(expiry),
                    "reminder_lead_days": lead,
                },
            ).status_code
            == 201
        )
    alerts = {
        a["permit_number"]: a
        for a in engine.get(
            f"/companies/{tenant_a.company_id}/permits/alerts", headers=bearer(tenant_a.tokens["VIEWER"])
        ).json()
    }
    assert alerts["P-DUE"]["alert"] == "due" and alerts["P-EXPIRED"]["alert"] == "expired"
    assert "P-LATER" not in alerts
    home = engine.get(f"/companies/{tenant_a.company_id}/home", headers=bearer(tenant_a.tokens["VIEWER"])).json()
    assert home["compliance"]["permit_alerts"] >= 2


# -- F16 dashboards --------------------------------------------------------------------------


def test_company_total_equals_sum_of_facilities_and_facility_view_is_scoped(engine, tenant_a, locked_project):
    from app.db import transaction
    from app.services import carbon
    from app.services.carbon import Filters

    with transaction(None) as tx:
        company = carbon.total(tx, Filters([locked_project]))
        by_fac = carbon.breakdown(tx, Filters([locked_project]), "facility")
        plant = carbon.total(tx, Filters([locked_project], facility_id=tenant_a.facility_id))
        drill = carbon.drill(tx, Filters([locked_project], facility_id=tenant_a.facility_id))
    assert len(by_fac) == 2
    assert sum(r["co2e_t"] for r in by_fac) == company
    assert plant == next(r["co2e_t"] for r in by_fac if r["key"] == tenant_a.facility_id)
    assert {d["facility"] for d in drill} == {"Plant 1"} and len(drill) == 2


def test_matched_month_delta():
    from app.services.carbon import matched_months_delta

    d = matched_months_delta(
        {"2026-01": Decimal(110), "2026-02": None, "2026-03": Decimal(90)},
        {"2025-01": Decimal(100), "2025-02": Decimal(100), "2025-03": Decimal(100), "2025-12": Decimal(1)},
    )
    assert d["months"] == 2 and d["pct"] == 0, "only months present in both years are compared"
    assert matched_months_delta({"2026-01": Decimal(1)}, {})["pct"] is None
