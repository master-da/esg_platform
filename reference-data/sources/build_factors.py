"""Build reference-data/lib/emission_factors.csv and factor_conversions.csv from official sources.

Inputs (download first; URLs and SHA-256 are listed in SOURCES below and in reference-data/README.md):
  --desnz2025  ghg-conversion-factors-2025-flat-format.xlsx
  --desnz2026  ghg-conversion-factors-2026-flat-format-revised.xlsx
  --epa        SupplyChainGHGEmissionFactors_v1.3.0_NAICS_CO2e_USD2022.csv

Rules (PRD Feature 6, "no fabrication"):
  * every value is copied from the published source; nothing is estimated or defaulted
  * a mapped row that is blank in the source produces NO factor (the leaf stays UNKNOWN)
  * conversion multipliers are ratios of the SAME fuel's published factors in two units, so a
    converted calculation reproduces the published factor for that unit exactly
  * manual_factors.csv holds hand-transcribed official values (e.g. Bangladesh grid factor) with citations

Run: python -I reference-data/sources/build_factors.py --desnz2025 ... --desnz2026 ... --epa ...
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
LIB = HERE.parent / "lib"

UOM = {"litres": "L", "tonnes": "t", "cubic metres": "m3", "kWh (Net CV)": "kWh", "kg": "kg",
       "passenger.km": "pkm", "km": "km", "tonne.km": "tkm", "Room per night": "room_night"}
GHG_COL = {"kg CO2e": "co2e", "kg CO2e of CO2 per unit": "co2e_co2", "kg CO2e of CH4 per unit": "co2e_ch4",
           "kg CO2e of N2O per unit": "co2e_n2o"}

SOURCES = {
    "desnz2025": {
        "name": "UK DESNZ GHG Conversion Factors for Company Reporting 2025 (flat file v1, 10 Jun 2025)",
        "url": "https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2025",
        "file_url": "https://assets.publishing.service.gov.uk/media/6846b6ea57f3515d9611f0dd/ghg-conversion-factors-2025-flat-format.xlsx",
        "sha256": "8bfdb45b81ec4a88e3bdf4584637330f62e6bd09ce1940e654c5d7b7f736de94",
        "year": 2025, "valid_from": "2025-01-01", "valid_to": "2025-12-31", "version": 1,
    },
    "desnz2026": {
        "name": "UK DESNZ GHG Conversion Factors for Company Reporting 2026 (flat file, revised Jul 2026)",
        "url": "https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2026",
        "file_url": "https://assets.publishing.service.gov.uk/media/6a6c9748862aaf18d9c62ac9/ghg-conversion-factors-2026-flat-format-revised.xlsx",
        "sha256": "a9a455ab396dae226d510c7be6233748416d490c41a5d20f3dc7a0c45feecd5e",
        "year": 2026, "valid_from": "2026-01-01", "valid_to": "2026-12-31", "version": 2,
    },
    "epa": {
        "name": "US EPA Supply Chain GHG Emission Factors v1.3 by NAICS-6 (2022 USD, purchaser price, with margins)",
        "url": "https://doi.org/10.23719/1531143",
        "file_url": "https://pasteur.epa.gov/uploads/10.23719/1531143/SupplyChainGHGEmissionFactors_v1.3.0_NAICS_CO2e_USD2022.csv",
        "sha256": "6025a14c4fe16675735efba1683030e4d36011e3973f3c4261f8231aa69c002f",
        "year": 2022,
        # Applicability window for spend data is a methodological choice pending domain-advisor sign-off.
        "valid_from": "2022-01-01", "valid_to": "2026-12-31", "version": 1,
    },
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dec(v) -> Decimal | None:
    if v is None or v == "":
        return None
    return Decimal(str(v))


def load_desnz(path: Path) -> tuple[dict[tuple, dict], dict[str, tuple]]:
    """Returns ({(L1, L2, L3, L4, column_text, uom): {co2e, co2e_co2, ..., "_id": base_id}},
                {base_id: label key}). The base ID (row ID without its per-gas suffix) is stable
    across annual releases even when DESNZ renames a row."""
    rows = list(load_workbook(path, read_only=True, data_only=True)["Factors by Category"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r and r[0] == "ID")
    out: dict[tuple, dict] = {}
    by_id: dict[str, tuple] = {}
    for r in rows[hdr + 1:]:
        if not r[0] or r[0] == "END" or r[8] not in GHG_COL:
            continue
        key = tuple("" if x is None else str(x) for x in (r[2], r[3], r[4], r[5], r[6], r[7]))
        base_id = str(r[0]).rsplit("_", 1)[0]
        entry = out.setdefault(key, {"_id": base_id})
        entry[GHG_COL[r[8]]] = dec(r[9])
        by_id[base_id] = key
    return out, by_id


def load_epa(path: Path) -> dict[str, dict]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return {row["2017 NAICS Code"]: row for row in csv.DictReader(fh)}


def read_map() -> list[dict]:
    with (HERE / "factor_map.csv").open(newline="") as fh:
        return list(csv.DictReader(line for line in fh if not line.startswith("#")))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--desnz2025", type=Path, required=True)
    ap.add_argument("--desnz2026", type=Path, required=True)
    ap.add_argument("--epa", type=Path, required=True)
    args = ap.parse_args()

    for key, path in (("desnz2025", args.desnz2025), ("desnz2026", args.desnz2026), ("epa", args.epa)):
        actual = sha256(path)
        if actual != SOURCES[key]["sha256"]:
            raise SystemExit(f"{path.name}: sha256 {actual} does not match the pinned source file")

    desnz, desnz_ids = {}, {}
    for k, p in (("desnz2025", args.desnz2025), ("desnz2026", args.desnz2026)):
        desnz[k], desnz_ids[k] = load_desnz(p)
    epa = load_epa(args.epa)
    factors: list[dict] = []
    conversions: list[dict] = []
    skipped: list[str] = []

    for m in read_map():
        leaf = m["leaf_code"]
        if m["source"] == "epa":
            row = epa.get(m["naics"])
            if row is None:
                raise SystemExit(f"{leaf}: NAICS {m['naics']} not in EPA file")
            src = SOURCES["epa"]
            value = dec(row["Supply Chain Emission Factors with Margins"])
            factors.append({
                "factor_code": f"{leaf}-GLOBAL", "taxonomy_code": leaf, "geography_code": "GLOBAL",
                "activity_unit": "USD", "co2e": value, "co2e_co2": "", "co2e_ch4": "", "co2e_n2o": "",
                "gwp_version": "AR5", "source_name": src["name"], "source_url": src["url"], "source_year": src["year"],
                "source_ref": f"NAICS {m['naics']} {row['2017 NAICS Title']}",
                "methodology": "Spend-based (EEIO): kg CO2e per 2022 USD, purchaser price, including margins",
                "quality_rating": "Spend-based estimate (grade C); US economy proxy",
                "valid_from": src["valid_from"], "valid_to": src["valid_to"], "version": src["version"],
                "notes": (m["note"] + "; " if m["note"] else "") +
                         "Spend in other currencies is converted with cited FX rates; no inflation adjustment to 2022 USD",
            })
            continue

        mapped_key = (m["level1"], m["level2"], m["level3"], m["level4"], m["column_text"])
        first = desnz["desnz2025"].get((*mapped_key, m["uom"]))
        if first is None:
            raise SystemExit(f"{leaf}: no desnz2025 row for {mapped_key} [{m['uom']}]")
        for skey in ("desnz2025", "desnz2026"):
            src = SOURCES[skey]
            level_key = mapped_key
            base = desnz[skey].get((*level_key, m["uom"]))
            if base is None:
                # Row renamed in this release: follow the stable DESNZ row ID.
                renamed = desnz_ids[skey].get(first["_id"])
                if renamed is None or renamed[-1] != m["uom"]:
                    raise SystemExit(f"{leaf}: no {skey} row for {level_key} [{m['uom']}] (id {first['_id']})")
                level_key = renamed[:-1]
                base = desnz[skey][renamed]
                skipped.append(f"{leaf} {skey}: row renamed to {' > '.join(x for x in level_key if x)} (same DESNZ id)")
            if base.get("co2e") is None:
                skipped.append(f"{leaf} {skey}: published value is blank -> no factor (UNKNOWN)")
                continue
            code = f"{leaf}-GLOBAL"
            factors.append({
                "factor_code": code, "taxonomy_code": leaf, "geography_code": "GLOBAL",
                "activity_unit": UOM[m["uom"]], "co2e": base["co2e"],
                "co2e_co2": base.get("co2e_co2") or "", "co2e_ch4": base.get("co2e_ch4") or "",
                "co2e_n2o": base.get("co2e_n2o") or "",
                "gwp_version": "AR5", "source_name": src["name"], "source_url": src["url"], "source_year": src["year"],
                "source_ref": " > ".join(x for x in level_key if x) + f" [{m['uom']}] (DESNZ id {base['_id']})",
                "methodology": "Activity-based: kg CO2e per unit of activity (DESNZ)",
                "quality_rating": "International default (UK DESNZ), applied globally; no country-specific value available",
                "valid_from": src["valid_from"], "valid_to": src["valid_to"], "version": src["version"],
                "notes": m["note"],
            })
            for other in filter(None, m["conversions"].split(";")):
                alt = desnz[skey].get((*level_key, other))
                if not alt or alt.get("co2e") in (None, Decimal(0)):
                    skipped.append(f"{leaf} {skey}: no published {other} factor -> no {other} conversion")
                    continue
                conversions.append({
                    "factor_code": code, "version": src["version"], "from_unit": UOM[other], "to_unit": UOM[m["uom"]],
                    "multiplier": alt["co2e"] / base["co2e"],
                    "source_name": src["name"], "source_url": src["url"],
                    "note": f"Ratio of the published {other} and {m['uom']} factors for this fuel "
                            f"(implied {'energy content' if 'kWh' in other else 'density'})",
                })

    with (HERE / "manual_factors.csv").open(newline="") as fh:
        for row in csv.DictReader(line for line in fh if not line.startswith("#")):
            factors.append(row)

    fields = ["factor_code", "taxonomy_code", "geography_code", "activity_unit", "co2e", "co2e_co2", "co2e_ch4",
              "co2e_n2o", "gwp_version", "source_name", "source_url", "source_year", "source_ref", "methodology",
              "quality_rating", "valid_from", "valid_to", "version", "notes"]
    with (LIB / "emission_factors.csv").open("w", newline="") as fh:
        fh.write("# Generated by sources/build_factors.py from pinned official sources. Do not edit by hand.\n")
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for f in factors:
            w.writerow({k: ("" if f.get(k) is None else f.get(k)) for k in fields})
    with (LIB / "factor_conversions.csv").open("w", newline="") as fh:
        fh.write("# Generated by sources/build_factors.py. Multipliers convert from_unit -> to_unit (the factor's unit).\n")
        w = csv.DictWriter(fh, fieldnames=["factor_code", "version", "from_unit", "to_unit", "multiplier",
                                           "source_name", "source_url", "note"])
        w.writeheader()
        w.writerows(conversions)
    print(f"{len(factors)} factors, {len(conversions)} conversions")
    for s in skipped:
        print("  note:", s)


if __name__ == "__main__":
    main()
